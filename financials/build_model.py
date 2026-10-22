from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.chart import LineChart, BarChart, Reference

F = 'Arial'
def font(**k): return Font(name=F, size=k.pop('size', 10), **k)
BLUE = font(color='0000FF'); BLACK = font(); GREEN = font(color='008000'); BOLD = font(bold=True)
TITLE = font(bold=True, size=14, color='0B2340'); HEAD = font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='0B2340'); SUB = PatternFill('solid', fgColor='EEF2F7'); YEL = PatternFill('solid', fgColor='FFFF00')
USD = '$#,##0;($#,##0);-'; USD2 = '$#,##0.00;($#,##0.00);-'; PCT = '0.0%;(0.0%);-'; PCT2 = '0.00%;(0.00%);-'; NUM = '#,##0;(#,##0);-'; NUM1 = '#,##0.0;(#,##0.0);-'; X = '0.0x'
thin = Side(style='thin', color='BBBBBB')

wb = Workbook()
A = wb.active; A.title = 'Assumptions'
def put(ws, ref, v, f=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]; c.value = v; c.font = f
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if note: c.comment = Comment(note, 'Model')
    return c
def header(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t); c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')

put(A, 'A1', 'InjuryCase — 5-year financial model: assumptions', TITLE)
put(A, 'A2', 'Blue = input you can change · yellow = key assumption · black = formula · green = link to another sheet. Monthly sheets and the Summary recalculate from these cells.', font(italic=True, color='555555'))

# ---- market segments ----
put(A, 'A4', 'Market segments (US practices that treat injury patients)', BOLD)
header(A, 5, ['Segment', 'US practices', '% treating PI (est.)', 'Addressable practices', 'Launch month', 'Monthly capture — Conservative', 'Monthly capture — Base', 'Monthly capture — Aggressive', 'Price multiplier (plan mix vs. chiro)', 'Source / basis'])
segs = [
    ('Chiropractic clinics', 66061, 0.35, 1, (0.0006, 0.0010, 0.0016), 1.00,
     'Practices: IBISWorld, Chiropractors in the US, 2026 business count (ibisworld.com/united-states/number-of-businesses/chiropractors/1559). PI share is an estimate; NBCE Practice Analysis: ~14% of a typical chiropractor\'s patients come after a motor-vehicle accident.'),
    ('Physical therapy clinics', 39000, 0.20, 7, (0.0004, 0.0007, 0.0012), 1.10,
     'Practices: ~38,000–40,000 outpatient PT clinics (APTA-based estimate, withorbital.com/data/largest-physical-therapy-chains-in-the-us). PI share estimated. Launch after PT-specific care-guide changes.'),
    ('Medical injury clinics (MD/DO, urgent care treating PI)', 6000, 0.50, 7, (0.0005, 0.0008, 0.0014), 1.15,
     'No census found — 6,000 is a planning estimate. PI share estimated.'),
    ('Orthopedic practices', 23142, 0.15, 19, (0.0003, 0.0005, 0.0009), 1.40,
     'Practices: Ciro, How many orthopedic practices are in the US (2024): 23,142 tracked. PI share estimated. Launch after a specialty module (procedures, surgical liens).'),
    ('Pain management practices', 5000, 0.35, 19, (0.0004, 0.0007, 0.0012), 1.40,
     'Planning estimate: Definitive Healthcare lists 900+ pain-management groups (Dec 2025), most with several sites. PI share estimated. Launch after the specialty module.'),
]
S0 = 6
for i, (name, n, pi, launch, cap, mult, src) in enumerate(segs):
    r = S0 + i
    put(A, f'A{r}', name)
    put(A, f'B{r}', n, BLUE, NUM)
    put(A, f'C{r}', pi, BLUE, PCT, YEL)
    put(A, f'D{r}', f'=B{r}*C{r}', BLACK, NUM)
    put(A, f'E{r}', launch, BLUE, NUM)
    for j, v in enumerate(cap): put(A, f'{"FGH"[j]}{r}', v, BLUE, PCT2, YEL)
    put(A, f'I{r}', mult, BLUE, '0.00x')
    put(A, f'J{r}', src, font(size=9, color='555555'))
SE = S0 + len(segs) - 1
put(A, f'A{SE+1}', 'Total', BOLD); put(A, f'B{SE+1}', f'=SUM(B{S0}:B{SE})', BOLD, NUM); put(A, f'D{SE+1}', f'=SUM(D{S0}:D{SE})', BOLD, NUM)
put(A, f'J{SE+1}', 'Monthly capture = share of a segment\'s addressable practices that become paying customers each month once marketing is fully ramped.', font(size=9, color='555555'))

# ---- pricing ----
r = SE + 3
put(A, f'A{r}', 'Pricing (current plans)', BOLD); r += 1
header(A, r, ['Plan', 'Monthly price', 'Plan mix', '', '', '', '', '', '', 'Notes']); r += 1
P0 = r
for name, price, mix in [('Silver', 39, 0.40), ('Gold', 59, 0.45), ('Platinum', 79, 0.15)]:
    put(A, f'A{r}', name); put(A, f'B{r}', price, BLUE, USD); put(A, f'C{r}', mix, BLUE, PCT, YEL); r += 1
PE = r - 1
put(A, f'J{P0}', 'Prices from app/saas.php (PLANS). Plan mix is an estimate for chiropractic clinics; segment price multipliers lift it for larger practices.', font(size=9, color='555555'))
put(A, f'A{r}', 'Plan mix check (must be 100%)'); put(A, f'C{r}', f'=SUM(C{P0}:C{PE})', BLACK, PCT); r += 1
ANN = r; put(A, f'A{r}', 'Customers on annual billing'); put(A, f'B{r}', 0.30, BLUE, PCT, YEL); put(A, f'J{r}', 'Annual plans are charged 10 months for 12 (ANNUAL_MONTHS_CHARGED).', font(size=9, color='555555')); r += 1
ANM = r; put(A, f'A{r}', 'Months charged on an annual plan'); put(A, f'B{r}', 10, BLUE, NUM); r += 1
LMULT = r; put(A, f'A{r}', 'List-price multiplier (1.0 = today\'s prices)'); put(A, f'B{r}', 1.0, BLUE, '0.00x', YEL); put(A, f'J{r}', 'Pricing lever: try 1.5 or 2.0 to test higher prices.', font(size=9, color='555555')); r += 1
ARPU = r; put(A, f'A{r}', 'Blended revenue per chiro customer per month', BOLD); put(A, f'B{r}', f'=SUMPRODUCT(B{P0}:B{PE},C{P0}:C{PE})*((1-B{ANN})+B{ANN}*B{ANM}/12)*B{LMULT}', BOLD, USD2); r += 2

# ---- scenario levers ----
put(A, f'A{r}', 'Scenario levers', BOLD); r += 1
header(A, r, ['Lever', '', 'Conservative', 'Base', 'Aggressive', '', '', '', '', 'Notes']); r += 1
LEV = {}
for key, label, vals, fmt, note in [
    ('churn', 'Monthly customer churn', (0.035, 0.025, 0.018), PCT, 'Small-practice SaaS commonly loses 1.5–4% of customers a month.'),
    ('cac', 'Acquisition cost per new paying customer', (600, 450, 350), USD, 'Ads, conferences, attorney and association referrals, sales time.'),
    ('inc', 'Annual price increase', (0.0, 0.03, 0.05), PCT, 'Applied at the start of each year after year 1.'),
    ('ramp', 'Months to reach full marketing ramp per segment', (15, 12, 9), NUM, 'New customers grow linearly from launch to the full capture rate.'),
]:
    put(A, f'A{r}', label)
    for j, v in enumerate(vals): put(A, f'{"CDE"[j]}{r}', v, BLUE, fmt, YEL)
    put(A, f'J{r}', note, font(size=9, color='555555')); LEV[key] = r; r += 1
r += 1

# ---- costs ----
put(A, f'A{r}', 'Costs (all scenarios)', BOLD); r += 1
header(A, r, ['Item', 'Value', '', '', '', '', '', '', '', 'Notes']); r += 1
C = {}
for key, label, v, fmt, note in [
    ('stripe_pct', 'Card processing: % of revenue', 0.029, PCT, 'Stripe standard US card pricing.'),
    ('stripe_fix', 'Card processing: fee per charge', 0.30, USD2, 'Stripe standard US card pricing.'),
    ('host_base', 'Hosting: fixed per month', 250, USD, 'AWS stack (deploy/aws): ~$40/mo minimum; budgeted higher for Multi-AZ database, backups and monitoring.'),
    ('host_cust', 'Hosting: per customer per month', 0.75, USD2, 'Storage, email (SES) and compute growth.'),
    ('sup_ratio', 'Customers per support person', 400, NUM, 'Estimate.'),
    ('sup_sal', 'Support salary, fully loaded (per year)', 55000, USD, 'Estimate.'),
    ('eng_sal', 'Engineer, fully loaded (per year)', 140000, USD, 'Estimate; can be contractors.'),
    ('mkt_fixed', 'Marketing: fixed per month', 1500, USD, 'Website, content, tools.'),
    ('legal', 'Legal review: one-time in month 1', 20000, USD, 'Attorney review of Terms, BAA and state rules (flagged as needed in the app).'),
    ('soc2_amt', 'SOC 2 audit: one-time', 30000, USD, 'Often required by larger groups.'),
    ('soc2_m', 'SOC 2 audit: month', 25, NUM, ''),
    ('comp', 'Compliance and legal: per month', 750, USD, 'Ongoing counsel, HIPAA program.'),
    ('ins', 'Insurance (cyber, E&O): per month', 450, USD, 'Estimate.'),
    ('gna', 'Other G&A (software, accounting): per month', 600, USD, 'Estimate.'),
    ('cash0', 'Starting cash', 50000, USD, 'Change to your actual funds.'),
]:
    put(A, f'A{r}', label); put(A, f'B{r}', v, BLUE, fmt, YEL if key in ('cash0', 'sup_ratio') else None)
    put(A, f'J{r}', note, font(size=9, color='555555')); C[key] = r; r += 1
r += 1
header(A, r, ['Headcount and pay by year', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5', '', '', '', 'Notes']); r += 1
ENG = r; put(A, f'A{r}', 'Engineers (FTE)')
for j, v in enumerate([0.25, 0.5, 1.0, 1.5, 2.0]): put(A, f'{L(2+j)}{r}', v, BLUE, '0.00')
put(A, f'J{r}', 'Lean build: founder plus contractors; grows with the specialty modules.', font(size=9, color='555555')); r += 1
FND = r; put(A, f'A{r}', 'Founder salary (per year)')
for j, v in enumerate([0, 0, 60000, 90000, 120000]): put(A, f'{L(2+j)}{r}', v, BLUE, USD)
r += 1

A.column_dimensions['A'].width = 46
for col in 'BCDEFGHI': A.column_dimensions[col].width = 15
A.column_dimensions['J'].width = 90
A.row_dimensions[5].height = 42
A.freeze_panes = 'B6'

# ---- monthly sheets ----
M = 60; MC0 = 3  # first month column C
lastc = L(MC0 + M - 1)
def mcol(m): return L(MC0 + m - 1)
SC = {'Conservative': 'C', 'Base': 'D', 'Aggressive': 'E'}
CAPC = {'Conservative': 'F', 'Base': 'G', 'Aggressive': 'H'}
rows = {}
for scen in ['Conservative', 'Base', 'Aggressive']:
    ws = wb.create_sheet(f'Monthly {scen}')
    sc = SC[scen]; cc = CAPC[scen]
    put(ws, 'A1', f'Monthly projection — {scen} scenario', TITLE)
    put(ws, 'A2', 'All values are formulas driven by the Assumptions sheet.', font(italic=True, color='555555'))
    R = {}
    rr = 4
    def row(key, label, bold=False):
        global rr
        R[key] = rr; put(ws, f'A{rr}', label, BOLD if bold else BLACK); rr += 1
        return R[key]
    row('m', 'Month', True); row('y', 'Year', True); row('pf', 'Price factor (increases × list multiplier)')
    rr += 1; put(ws, f'A{rr}', 'Marketing ramp by segment', BOLD); rr += 1
    for i in range(len(segs)): row(f'ramp{i}', segs[i][0])
    rr += 1; put(ws, f'A{rr}', 'New paying customers', BOLD); rr += 1
    for i in range(len(segs)): row(f'new{i}', segs[i][0])
    rr += 1; put(ws, f'A{rr}', 'Active customers (end of month)', BOLD); rr += 1
    for i in range(len(segs)): row(f'act{i}', segs[i][0])
    rr += 1
    row('newt', 'Total new customers', True); row('churn', 'Customers lost (churn)'); row('act', 'Total active customers', True)
    rr += 1
    row('mrr', 'MRR ($)', True); row('arr', 'ARR ($)', True); row('rev', 'Revenue ($)', True)
    rr += 1; put(ws, f'A{rr}', 'Cost of revenue', BOLD); rr += 1
    row('stripe', 'Card processing'); row('host', 'Hosting'); row('sup', 'Customer support'); row('cogs', 'Total cost of revenue', True)
    row('gp', 'Gross profit', True); row('gm', 'Gross margin %')
    rr += 1; put(ws, f'A{rr}', 'Operating expenses', BOLD); rr += 1
    row('eng', 'Engineering'); row('fnd', 'Founder salary'); row('mkt', 'Sales & marketing'); row('leg', 'Legal & compliance'); row('ins', 'Insurance'); row('gna', 'Other G&A'); row('opex', 'Total operating expenses', True)
    rr += 1
    row('ebitda', 'EBITDA', True); row('cash', 'Cash balance (end of month)', True)
    rows[scen] = R
    for m in range(1, M + 1):
        c = mcol(m); p = mcol(m - 1) if m > 1 else None
        put(ws, f'{c}{R["m"]}', m, BOLD, NUM)
        put(ws, f'{c}{R["y"]}', f'=INT(({c}{R["m"]}-1)/12)+1', BOLD, NUM)
        put(ws, f'{c}{R["pf"]}', f"=(1+Assumptions!${sc}${LEV['inc']})^({c}{R['y']}-1)*Assumptions!$B${LMULT}", BLACK, '0.000')
        for i in range(len(segs)):
            ar = S0 + i
            put(ws, f'{c}{R[f"ramp{i}"]}', f"=IF({c}{R['m']}<Assumptions!$E${ar},0,MIN(1,({c}{R['m']}-Assumptions!$E${ar}+1)/Assumptions!${sc}${LEV['ramp']}))", BLACK, '0%')
            put(ws, f'{c}{R[f"new{i}"]}', f"=Assumptions!$D${ar}*Assumptions!${cc}${ar}*{c}{R[f'ramp{i}']}", BLACK, NUM1)
            prev = f"{p}{R[f'act{i}']}*(1-Assumptions!${sc}${LEV['churn']})+" if p else ''
            put(ws, f'{c}{R[f"act{i}"]}', f"={prev}{c}{R[f'new{i}']}", BLACK, NUM1)
        put(ws, f'{c}{R["newt"]}', f'=SUM({c}{R["new0"]}:{c}{R[f"new{len(segs)-1}"]})', BOLD, NUM1)
        put(ws, f'{c}{R["churn"]}', (f"={p}{R['act']}*Assumptions!${sc}${LEV['churn']}" if p else 0), BLACK, NUM1)
        put(ws, f'{c}{R["act"]}', f'=SUM({c}{R["act0"]}:{c}{R[f"act{len(segs)-1}"]})', BOLD, NUM1)
        put(ws, f'{c}{R["mrr"]}', f"=SUMPRODUCT({c}{R['act0']}:{c}{R[f'act{len(segs)-1}']},Assumptions!$I${S0}:$I${SE})*Assumptions!$B${ARPU}/Assumptions!$B${LMULT}*{c}{R['pf']}", BOLD, USD)
        put(ws, f'{c}{R["arr"]}', f'={c}{R["mrr"]}*12', BOLD, USD)
        put(ws, f'{c}{R["rev"]}', f'={c}{R["mrr"]}', BOLD, USD)
        put(ws, f'{c}{R["stripe"]}', f"={c}{R['rev']}*Assumptions!$B${C['stripe_pct']}+{c}{R['act']}*((1-Assumptions!$B${ANN})+Assumptions!$B${ANN}/12)*Assumptions!$B${C['stripe_fix']}", BLACK, USD)
        put(ws, f'{c}{R["host"]}', f"=Assumptions!$B${C['host_base']}+Assumptions!$B${C['host_cust']}*{c}{R['act']}", BLACK, USD)
        put(ws, f'{c}{R["sup"]}', f"={c}{R['act']}/Assumptions!$B${C['sup_ratio']}*Assumptions!$B${C['sup_sal']}/12", BLACK, USD)
        put(ws, f'{c}{R["cogs"]}', f'=SUM({c}{R["stripe"]}:{c}{R["sup"]})', BOLD, USD)
        put(ws, f'{c}{R["gp"]}', f'={c}{R["rev"]}-{c}{R["cogs"]}', BOLD, USD)
        put(ws, f'{c}{R["gm"]}', f'=IF({c}{R["rev"]}=0,0,{c}{R["gp"]}/{c}{R["rev"]})', BLACK, PCT)
        put(ws, f'{c}{R["eng"]}', f"=INDEX(Assumptions!$B${ENG}:$F${ENG},{c}{R['y']})*Assumptions!$B${C['eng_sal']}/12", BLACK, USD)
        put(ws, f'{c}{R["fnd"]}', f"=INDEX(Assumptions!$B${FND}:$F${FND},{c}{R['y']})/12", BLACK, USD)
        put(ws, f'{c}{R["mkt"]}', f"={c}{R['newt']}*Assumptions!${sc}${LEV['cac']}+Assumptions!$B${C['mkt_fixed']}", BLACK, USD)
        put(ws, f'{c}{R["leg"]}', f"=Assumptions!$B${C['comp']}+IF({c}{R['m']}=1,Assumptions!$B${C['legal']},0)+IF({c}{R['m']}=Assumptions!$B${C['soc2_m']},Assumptions!$B${C['soc2_amt']},0)", BLACK, USD)
        put(ws, f'{c}{R["ins"]}', f"=Assumptions!$B${C['ins']}", GREEN, USD)
        put(ws, f'{c}{R["gna"]}', f"=Assumptions!$B${C['gna']}", GREEN, USD)
        put(ws, f'{c}{R["opex"]}', f'=SUM({c}{R["eng"]}:{c}{R["gna"]})', BOLD, USD)
        put(ws, f'{c}{R["ebitda"]}', f'={c}{R["gp"]}-{c}{R["opex"]}', BOLD, USD)
        put(ws, f'{c}{R["cash"]}', (f'={p}{R["cash"]}+{c}{R["ebitda"]}' if p else f"=Assumptions!$B${C['cash0']}+{c}{R['ebitda']}"), BOLD, USD)
        ws.column_dimensions[c].width = 11
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 2
    ws.freeze_panes = 'C6'

# ---- summary ----
Sm = wb.create_sheet('Summary', 0)
put(Sm, 'A1', 'InjuryCase — 5-year projection summary', TITLE)
put(Sm, 'A2', 'Planning model, not a forecast. Market counts and capture rates are estimates — see Assumptions and Sources. Edit blue cells on Assumptions.', font(italic=True, color='555555'))
r = 4
yc = {y: L(1 + y) for y in range(1, 6)}
blocks = {}
metrics = [('cust', 'Customers (end of year)', NUM), ('arr', 'ARR, end of year ($)', USD), ('rev', 'Revenue ($)', USD), ('gp', 'Gross profit ($)', USD), ('gm', 'Gross margin %', PCT),
           ('opex', 'Operating expenses ($)', USD), ('ebitda', 'EBITDA ($)', USD), ('cash', 'Cash, end of year ($)', USD)]
for scen in ['Conservative', 'Base', 'Aggressive']:
    R = rows[scen]; sh = f"'Monthly {scen}'"
    put(Sm, f'A{r}', f'{scen} scenario', BOLD); Sm[f'A{r}'].fill = SUB
    header(Sm, r + 1, ['($, US)', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5'])
    rr = r + 2; blocks[scen] = {}
    for key, label, fmt in metrics:
        put(Sm, f'A{rr}', label); blocks[scen][key] = rr
        for y in range(1, 6):
            c = yc[y]; end = f'{mcol(12*y)}'
            yrng = f"{sh}!$C${R['y']}:${lastc}${R['y']}"
            if key == 'cust': f = f"={sh}!{end}{R['act']}"
            elif key == 'arr': f = f"={sh}!{end}{R['arr']}"
            elif key == 'cash': f = f"={sh}!{end}{R['cash']}"
            elif key == 'gm': f = f'=IF({c}{blocks[scen]["rev"]}=0,0,{c}{blocks[scen]["gp"]}/{c}{blocks[scen]["rev"]})'
            else: f = f"=SUMPRODUCT(({yrng}={y})*{sh}!$C${R[key]}:${lastc}${R[key]})"
            put(Sm, f'{c}{rr}', f, GREEN if key in ('cust', 'arr', 'cash') else BLACK, fmt)
        rr += 1
    lt = rr
    put(Sm, f'A{rr}', 'Lowest cash balance in 5 years ($)'); put(Sm, f'B{rr}', f"=MIN({sh}!$C${R['cash']}:${lastc}${R['cash']})", BLACK, USD); rr += 1
    put(Sm, f'A{rr}', 'Funding needed to stay above $0 ($)', BOLD); put(Sm, f'B{rr}', f'=MAX(0,-B{lt})', BOLD, USD); rr += 1
    put(Sm, f'A{rr}', 'First month with positive EBITDA'); put(Sm, f'B{rr}', f"=IFERROR(INDEX({sh}!$C${R['m']}:${lastc}${R['m']},MATCH(TRUE,INDEX({sh}!$C${R['ebitda']}:${lastc}${R['ebitda']}>0,0),0)),\"Not within 5 years\")", BLACK, NUM); rr += 1
    sc = SC[scen]
    put(Sm, f'A{rr}', 'Customer lifetime value ÷ acquisition cost (year 5)'); put(Sm, f'B{rr}', f"=IF({sh}!{mcol(60)}{R['act']}=0,0,{sh}!{mcol(60)}{R['mrr']}/{sh}!{mcol(60)}{R['act']}*{yc[5]}{blocks[scen]['gm']}/Assumptions!${sc}${LEV['churn']}/Assumptions!${sc}${LEV['cac']})", BLACK, X); rr += 1
    put(Sm, f'A{rr}', 'Share of addressable practices reached (year 5)'); put(Sm, f'B{rr}', f"={yc[5]}{blocks[scen]['cust']}/Assumptions!$D${SE+1}", BLACK, PCT2); rr += 1
    r = rr + 2

# chart data table
put(Sm, f'A{r}', 'ARR by scenario, end of year ($)', BOLD); ct = r + 1
header(Sm, ct, ['Scenario', 'Year 1', 'Year 2', 'Year 3', 'Year 4', 'Year 5'])
for i, scen in enumerate(['Conservative', 'Base', 'Aggressive']):
    put(Sm, f'A{ct+1+i}', scen)
    for y in range(1, 6): put(Sm, f'{yc[y]}{ct+1+i}', f'={yc[y]}{blocks[scen]["arr"]}', BLACK, USD)
ch = LineChart(); ch.title = 'ARR by scenario'; ch.y_axis.title = 'ARR ($)'; ch.x_axis.title = 'Year'; ch.height = 8; ch.width = 16
data = Reference(Sm, min_col=1, max_col=6, min_row=ct + 1, max_row=ct + 3)
ch.add_data(data, from_rows=True, titles_from_data=True)
ch.set_categories(Reference(Sm, min_col=2, max_col=6, min_row=ct))
Sm.add_chart(ch, 'H4')
cb = BarChart(); cb.title = 'Customers at year end (Base)'; cb.height = 8; cb.width = 16
cb.add_data(Reference(Sm, min_col=1, max_col=6, min_row=blocks['Base']['cust'], max_row=blocks['Base']['cust']), from_rows=True, titles_from_data=True)
cb.set_categories(Reference(Sm, min_col=2, max_col=6, min_row=ct)); cb.legend = None
Sm.add_chart(cb, 'H22')
Sm.column_dimensions['A'].width = 48
for y in range(1, 6): Sm.column_dimensions[yc[y]].width = 15
for scen in blocks:
    for c in 'ABCDEF': Sm[f'{c}{blocks[scen]["cust"]-1}'].alignment = Alignment(horizontal='center')

# ---- sources ----
So = wb.create_sheet('Sources & notes')
put(So, 'A1', 'Sources and notes', TITLE)
notes = [
    ('Chiropractic practices', 'IBISWorld, Chiropractors in the US — Number of Businesses (66,061 in 2026, updated March 2026): https://www.ibisworld.com/united-states/number-of-businesses/chiropractors/1559/'),
    ('Chiropractic PI share', 'NBCE Practice Analysis of Chiropractic (2015): about 14% of a typical chiropractor\'s patients seek care after a motor-vehicle accident: https://www.nbce.org/wp-content/uploads/chapter_08.pdf. The 35% of clinics treating PI is an estimate.'),
    ('Physical therapy clinics', 'About 38,000–40,000 outpatient PT clinics (APTA-based, vendor estimate): https://www.withorbital.com/data/largest-physical-therapy-chains-in-the-us/'),
    ('Orthopedic practices', 'Ciro, How many orthopedic practices are in the US (2024), 23,142 tracked: https://www.ciro.io/blogs/how-many-orthopedic-practices-are-in-the-us-2024'),
    ('Pain management', 'Definitive Healthcare, largest pain-management physician groups (900+ groups, Dec 2025): https://www.definitivehc.com/resources/healthcare-insights/largest-pain-management-physician-groups. The 5,000-location figure is an estimate.'),
    ('Medical injury clinics', 'No census found; 6,000 is a planning estimate.'),
    ('Prices', 'Silver $39, Gold $59, Platinum $79 per month; annual plans charged 10 months (app/saas.php).'),
    ('Product readiness', 'Today the workflow is built around the primary treating provider (chiropractic first). PT and medical injury clinics need light changes (care guide wording, visit types), so they launch in month 7. Orthopedic and pain-management practices are mostly referral recipients with procedures, surgical liens and higher bills, so they launch in month 19 after a specialty module. Change the launch months on Assumptions.'),
    ('What the model leaves out', 'Free-trial timing (customers are counted when they pay), taxes, depreciation, financing costs and working capital. Revenue is recognized monthly, including annual plans.'),
    ('Status', 'Planning model, not a forecast or investment advice. Replace the estimates with your own data (trial-to-paid rate, churn and acquisition cost) as soon as you have it.'),
]
for i, (k, v) in enumerate(notes):
    put(So, f'A{3+i}', k, BOLD); c = put(So, f'B{3+i}', v); c.alignment = Alignment(wrap_text=True, vertical='top'); So[f'A{3+i}'].alignment = Alignment(vertical='top')
So.column_dimensions['A'].width = 28; So.column_dimensions['B'].width = 120

out = '/home/user/CovChiro/financials/InjuryCase financial model.xlsx'
wb.save(out); print(out)
